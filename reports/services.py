import csv
import io
import calendar
from decimal import Decimal, InvalidOperation
from datetime import datetime, date
from django.db import models
from django.utils import timezone
from django.db.models import Sum
from django.http import HttpResponse
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from transactions.models import Transaction
from wallets.models import Account
from budgets.models import Category
from debts.models import Debt
from debts.services import get_debt_summary

def get_monthly_report_data(user, year, month):
    _, last_day = calendar.monthrange(year, month)
    start_date = date(year, month, 1)
    end_date = date(year, month, last_day)

    txs = Transaction.objects.filter(user=user, date__range=[start_date, end_date])
    income = txs.filter(transaction_type=Transaction.TYPE_INCOME).aggregate(tot=Sum('amount'))['tot'] or Decimal('0.00')
    expense = txs.filter(transaction_type=Transaction.TYPE_EXPENSE).aggregate(tot=Sum('amount'))['tot'] or Decimal('0.00')
    net = income - expense
    savings_rate = (net / income * Decimal('100.0')).quantize(Decimal('0.1')) if income > Decimal('0.00') else Decimal('0.00')

    # Categories breakdown
    cat_breakdown = []
    cat_items = txs.filter(transaction_type=Transaction.TYPE_EXPENSE).values(
        'category__name', 'category__color'
    ).annotate(total=Sum('amount')).order_by('-total')

    for c in cat_items:
        c_tot = c['total'] or Decimal('0.00')
        pct = (c_tot / expense * Decimal('100.0')).quantize(Decimal('0.1')) if expense > Decimal('0.00') else Decimal('0.00')
        cat_breakdown.append({
            'name': c['category__name'] or 'Uncategorized',
            'color': c['category__color'] or '#64748b',
            'total': c_tot,
            'percentage': pct,
        })

    # Account breakdown
    acc_items = txs.values('account__name').annotate(
        total_in=Sum('amount', filter=models.Q(transaction_type=Transaction.TYPE_INCOME) if hasattr(models, 'Q') else None)
    )

    return {
        'year': year,
        'month': month,
        'month_name': calendar.month_name[month],
        'income': income,
        'expense': expense,
        'net': net,
        'savings_rate': savings_rate,
        'categories': cat_breakdown,
        'transactions_count': txs.count(),
    }


def get_yearly_report_data(user, year):
    months_data = []
    total_income = Decimal('0.00')
    total_expense = Decimal('0.00')

    for m in range(1, 13):
        _, last_day = calendar.monthrange(year, m)
        start_date = date(year, m, 1)
        end_date = date(year, m, last_day)

        txs = Transaction.objects.filter(user=user, date__range=[start_date, end_date])
        inc = txs.filter(transaction_type=Transaction.TYPE_INCOME).aggregate(tot=Sum('amount'))['tot'] or Decimal('0.00')
        exp = txs.filter(transaction_type=Transaction.TYPE_EXPENSE).aggregate(tot=Sum('amount'))['tot'] or Decimal('0.00')
        net = inc - exp

        total_income += inc
        total_expense += exp

        months_data.append({
            'month': m,
            'month_name': calendar.month_name[m],
            'income': inc,
            'expense': exp,
            'net': net,
        })

    total_net = total_income - total_expense
    savings_rate = (total_net / total_income * Decimal('100.0')).quantize(Decimal('0.1')) if total_income > Decimal('0.00') else Decimal('0.00')

    return {
        'year': year,
        'months': months_data,
        'total_income': total_income,
        'total_expense': total_expense,
        'total_net': total_net,
        'savings_rate': savings_rate,
    }


# EXPORT SERVICES
def export_transactions_csv_response(user):
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = f'attachment; filename="transactions_{timezone.now().strftime("%Y%m%d")}.csv"'

    writer = csv.writer(response)
    writer.writerow(['Date', 'Type', 'Amount', 'Account', 'To Account', 'Category', 'Description', 'Notes'])

    txs = Transaction.objects.filter(user=user).select_related('account', 'to_account', 'category').order_by('-date')
    for tx in txs:
        writer.writerow([
            tx.date.strftime('%Y-%m-%d'),
            tx.transaction_type,
            str(tx.amount),
            tx.account.name,
            tx.to_account.name if tx.to_account else '',
            tx.category.name if tx.category else '',
            tx.description,
            tx.notes,
        ])
    return response


def export_transactions_excel_response(user):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Transactions"

    headers = ['Date', 'Type', 'Amount', 'Account', 'To Account', 'Category', 'Description', 'Notes']
    ws.append(headers)

    # Style header row
    header_fill = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font

    txs = Transaction.objects.filter(user=user).select_related('account', 'to_account', 'category').order_by('-date')
    for tx in txs:
        ws.append([
            tx.date.strftime('%Y-%m-%d'),
            tx.transaction_type,
            float(tx.amount),
            tx.account.name,
            tx.to_account.name if tx.to_account else '',
            tx.category.name if tx.category else '',
            tx.description,
            tx.notes,
        ])

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="transactions_{timezone.now().strftime("%Y%m%d")}.xlsx"'
    wb.save(response)
    return response


def export_debts_excel_response(user):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Debts"

    headers = ['Direction', 'Title', 'Creditor/Debtor', 'Principal', 'Interest Rate (%)', 'Remaining Balance', 'Due Date', 'Status']
    ws.append(headers)

    header_fill = PatternFill(start_color="6366F1", end_color="6366F1", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font

    debts = Debt.objects.filter(user=user).order_by('due_date')
    for d in debts:
        s = get_debt_summary(d)
        ws.append([
            d.get_debt_type_display(),
            d.title,
            d.creditor_or_debtor,
            float(d.principal),
            float(d.interest_rate),
            float(s['remaining_balance']),
            d.due_date.strftime('%Y-%m-%d'),
            s['status']
        ])

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="debts_{timezone.now().strftime("%Y%m%d")}.xlsx"'
    wb.save(response)
    return response


# CSV IMPORT SERVICES WITH PREVIEW
def parse_and_validate_csv_file(user, file):
    """
    Parses an uploaded CSV file and validates rows before committing.
    Expected headers (case-insensitive):
    Date, Type, Amount, Account, Category (optional), Description, Notes (optional)
    """
    decoded_file = file.read().decode('utf-8-sig').splitlines()
    reader = csv.DictReader(decoded_file)
    
    valid_rows = []
    invalid_rows = []

    user_accounts = {a.name.lower(): a for a in Account.objects.filter(user=user)}
    user_categories = {c.name.lower(): c for c in Category.objects.filter(user=user)}
    default_acc = Account.objects.filter(user=user, is_active=True).first()

    for idx, row in enumerate(reader, start=2):
        cleaned = {k.strip().lower(): v.strip() for k, v in row.items() if k}
        row_errors = []

        # Validate Date
        date_str = cleaned.get('date', '')
        parsed_date = None
        for fmt in ('%Y-%m-%d', '%m/%d/%Y', '%d/%m/%Y'):
            try:
                parsed_date = datetime.strptime(date_str, fmt).date()
                break
            except ValueError:
                pass
        if not parsed_date:
            row_errors.append(f"Invalid date format '{date_str}'. Expected YYYY-MM-DD or MM/DD/YYYY.")

        # Validate Type
        raw_type = cleaned.get('type', '').upper()
        if raw_type in ('INCOME', 'EXPENSE', 'TRANSFER'):
            tx_type = raw_type
        else:
            row_errors.append(f"Invalid type '{raw_type}'. Expected INCOME, EXPENSE, or TRANSFER.")
            tx_type = 'EXPENSE'

        # Validate Amount
        raw_amt = cleaned.get('amount', '').replace(',', '').replace('₱', '').replace('$', '')
        try:
            amt = Decimal(raw_amt)
            if amt <= Decimal('0.00'):
                row_errors.append("Amount must be greater than zero.")
        except (InvalidOperation, ValueError):
            row_errors.append(f"Invalid numeric amount '{raw_amt}'.")
            amt = Decimal('0.00')

        # Validate Account
        acc_name = cleaned.get('account', '')
        acc = user_accounts.get(acc_name.lower()) or default_acc
        if not acc:
            row_errors.append(f"Account '{acc_name}' does not exist and no fallback account found.")

        # Category
        cat_name = cleaned.get('category', '')
        cat = user_categories.get(cat_name.lower())

        desc = cleaned.get('description', '') or 'Imported Transaction'
        notes = cleaned.get('notes', '')

        row_payload = {
            'row_num': idx,
            'date': parsed_date.strftime('%Y-%m-%d') if parsed_date else date_str,
            'transaction_type': tx_type,
            'amount': str(amt),
            'account_id': acc.id if acc else None,
            'account_name': acc.name if acc else 'None',
            'category_id': cat.id if cat else None,
            'category_name': cat.name if cat else (cat_name or 'None'),
            'description': desc,
            'notes': notes,
        }

        if row_errors:
            row_payload['errors'] = row_errors
            invalid_rows.append(row_payload)
        else:
            valid_rows.append(row_payload)

    return {
        'valid_rows': valid_rows,
        'invalid_rows': invalid_rows,
        'total_rows': len(valid_rows) + len(invalid_rows),
    }


def commit_imported_transactions(user, valid_rows):
    """
    Saves validated rows to database.
    """
    created_count = 0
    for r in valid_rows:
        Transaction.objects.create(
            user=user,
            date=datetime.strptime(r['date'], '%Y-%m-%d').date(),
            transaction_type=r['transaction_type'],
            amount=Decimal(r['amount']),
            account_id=r['account_id'],
            category_id=r['category_id'],
            description=r['description'],
            notes=r.get('notes', '')
        )
        created_count += 1
    return created_count
